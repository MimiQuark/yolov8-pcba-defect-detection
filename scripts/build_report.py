from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
ANALYSIS_DIR = ROOT / 'data' / 'analysis'
TRAIN_DIR = ROOT / 'runs' / 'train' / 'pcba_yolov8n_960'
EVAL_DIR = ROOT / 'runs' / 'eval' / 'test_960'
OUTPUT = ROOT / 'reports' / 'PCBA_YOLOv8实验报告.docx'
ASSETS = ROOT / 'reports' / 'assets'


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding='utf-8'))


def set_east_asia_font(run, name: str = '微软雅黑') -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    set_east_asia_font(run, '宋体')


def add_table(document: Document, headers: list[str], rows: list[list[object]], caption: str | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    if caption:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(caption)
        run.font.size = Pt(9)
        run.italic = True
        set_east_asia_font(run, '宋体')


def add_image(document: Document, path: Path, caption: str, width_cm: float = 15.5) -> None:
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.font.size = Pt(9)
    run.italic = True
    set_east_asia_font(run, '宋体')



def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'none')
    bottom.set(qn('w:sz'), '0')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), 'auto')
    p_bdr.append(bottom)

def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def format_metric(value: object, percent: bool = False) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f'{number * 100:.2f}%' if percent else f'{number:.4f}'


def training_last_row(csv_path: Path) -> dict[str, str]:
    with csv_path.open('r', encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.DictReader(handle))
    return rows[-1] if rows else {}


def build_report() -> Path:
    summary_path = ANALYSIS_DIR / 'summary.json'
    metrics_path = EVAL_DIR / 'metrics.json'
    if not summary_path.exists():
        raise FileNotFoundError(f'缺少数据统计文件: {summary_path}')
    if not metrics_path.exists():
        raise FileNotFoundError(f'缺少测试集指标文件: {metrics_path}')

    summary = load_json(summary_path)
    metrics = load_json(metrics_path)
    overall = metrics['overall']
    per_class = metrics['per_class']
    train_config = load_json(TRAIN_DIR / 'training_config.json') if (TRAIN_DIR / 'training_config.json').exists() else {}
    last_epoch = training_last_row(TRAIN_DIR / 'results.csv') if (TRAIN_DIR / 'results.csv').exists() else {}

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    styles = document.styles
    styles['Normal'].font.name = '宋体'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    styles['Normal'].font.size = Pt(10.5)
    for style_name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = styles[style_name]
        style.font.name = '微软雅黑'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.color.rgb = None
    header = section.header.paragraphs[0]
    header.text = '基于 YOLOv8 的 PCBA 五类缺陷检测项目'
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    add_page_number(footer)
    footer.add_run(' 页')

    title = document.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('基于 YOLOv8 的 PCBA 五类缺陷检测项目')
    set_east_asia_font(run, '微软雅黑')
    remove_paragraph_borders(title)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('数据准备、模型训练、测试评估与网页演示')
    run.font.size = Pt(14)
    set_east_asia_font(run, '微软雅黑')
    info = document.add_table(rows=0, cols=2)
    info.style = 'Table Grid'
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in [
        ('项目路径', str(ROOT)),
        ('算法', 'Ultralytics YOLOv8n'),
        ('训练设备', '11th Gen Intel Core i5-11320H CPU'),
        ('数据规模', '600 张图片 / 4552 个目标框'),
        ('报告日期', datetime.now().strftime('%Y-%m-%d')),
    ]:
        cells = info.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], value)

    document.add_page_break()
    document.add_heading('摘要', level=1)
    document.add_paragraph(
        '本项目将 PCBA 的 Pascal VOC XML 标注转换为 YOLO 格式，基于 YOLOv8n 建立五类焊接缺陷检测模型，'
        '并完成训练、验证、测试集评估、图片与视频推理以及 Gradio 网页演示。系统输出缺陷类别、数量、位置和置信度，'
        '当任一缺陷被检测到时给出“不合格”结论。'
    )
    document.add_heading('1 项目背景与目标', level=1)
    document.add_paragraph(
        'PCBA 焊接缺陷具有目标小、类别间外观接近、背景纹理复杂等特点。传统人工目检效率低且一致性有限。'
        '本项目采用单阶段目标检测算法 YOLOv8，在保持 CPU 可运行的前提下完成从原始标注到可演示系统的完整闭环。'
    )
    document.add_paragraph('项目目标包括：建立可复现的数据转换流程；完成五类缺陷训练和测试；提供命令行与网页推理入口；生成可追溯的实验数据、模型和图表。')

    document.add_heading('2 数据集与预处理', level=1)
    document.add_paragraph(
        f"数据集包含 {summary['totals']['images']} 张 JPG 图片和 {summary['totals']['objects']} 个目标框，"
        '原始划分为训练集 350 张、验证集 100 张、测试集 150 张，图像尺寸均为 2448 × 2048。'
        '图片与 XML 一一对应，未发现未知类别、越界框或缺失标签。'
    )
    class_rows = []
    for name in summary['class_names']:
        class_rows.append(
            [
                summary['class_names_zh'][name],
                name,
                summary['totals']['class_objects'][name],
                summary['splits']['train']['class_objects'][name],
                summary['splits']['val']['class_objects'][name],
                summary['splits']['test']['class_objects'][name],
            ]
        )
    add_table(document, ['中文类别', '英文类别', '总数', '训练集', '验证集', '测试集'], class_rows, '表 2-1 五类 PCBA 缺陷分布')
    document.add_paragraph(
        f"数据集平均灰度范围为 {summary['brightness']['min']:.2f} 至 {summary['brightness']['max']:.2f}，"
        '均位于自动光照处理阈值 60 至 200 之间，因此未对现有图像执行亮度调整或 CLAHE，以避免破坏已经清洗的数据。'
    )
    add_image(document, ANALYSIS_DIR / 'samples_annotated.jpg', '图 2-1 PCBA 缺陷标注样例')
    add_image(document, ANALYSIS_DIR / 'class_distribution.png', '图 2-2 训练、验证和测试集的类别目标数量')
    add_image(document, ANALYSIS_DIR / 'brightness_hist.png', '图 2-3 图像平均灰度分布')
    document.add_heading('2.1 在线数据增广', level=2)
    document.add_paragraph(
        '训练阶段使用 HSV 色彩扰动、水平翻转、轻微旋转、平移、缩放和 Mosaic 拼接。'
        '验证集和测试集不执行增广，避免评估数据分布被改变。项目未复制训练图片，降低近重复样本带来的泄漏风险。'
    )

    document.add_heading('3 模型方法', level=1)
    document.add_paragraph(
        'YOLOv8 使用 CSPDarknet 风格的骨干网络、PAN-FPN 多尺度特征融合和 Anchor-Free 检测头。'
        '模型同时预测类别和边界框分布，适合本项目中小尺寸缺陷的定位与分类。'
    )
    document.add_paragraph('基础模型为 YOLOv8n，参数量约 3.0M，FLOPs 约 8.2G，适合在 CPU 环境完成演示和课程实验。')

    document.add_heading('4 实验设置', level=1)
    config_rows = [
        ['预训练模型', 'models/yolov8n.pt'],
        ['输入尺寸', '960 × 960'],
        ['Batch size', '2'],
        ['训练轮数', '40，patience=10'],
        ['优化器', 'AdamW，由 auto 模式自动确定'],
        ['学习率策略', '余弦退火，close_mosaic=10'],
        ['随机种子', '42'],
        ['设备', 'CPU'],
        ['精度', 'FP32，amp=false'],
    ]
    add_table(document, ['参数', '设置'], config_rows, '表 4-1 正式训练配置')
    document.add_paragraph('训练使用训练集更新参数，验证集用于每轮评估和最佳权重选择，测试集只在模型确定后执行最终评估。')

    document.add_heading('5 训练过程', level=1)
    add_image(document, TRAIN_DIR / 'results.png', '图 5-1 训练损失和验证指标曲线', width_cm=16.5)
    if last_epoch:
        selected = [
            ['train/box_loss', last_epoch.get('train/box_loss', '-')],
            ['train/cls_loss', last_epoch.get('train/cls_loss', '-')],
            ['train/dfl_loss', last_epoch.get('train/dfl_loss', '-')],
            ['metrics/precision(B)', last_epoch.get('metrics/precision(B)', '-')],
            ['metrics/recall(B)', last_epoch.get('metrics/recall(B)', '-')],
            ['metrics/mAP50(B)', last_epoch.get('metrics/mAP50(B)', '-')],
            ['metrics/mAP50-95(B)', last_epoch.get('metrics/mAP50-95(B)', '-')],
        ]
        add_table(document, ['指标', '最后一轮数值'], selected, '表 5-1 最后一轮训练与验证指标')
    add_image(document, TRAIN_DIR / 'val_batch0_pred.jpg', '图 5-2 验证集预测示例')
    document.add_paragraph('训练在第 36 轮触发早停，最佳验证 mAP50-95 出现在第 26 轮，因此最终模型使用第 26 轮保存的权重。')
    if train_config:
        document.add_paragraph(
            f"训练配置已保存到 {TRAIN_DIR / 'training_config.json'}，最佳权重由训练过程自动复制到 models 目录。"
        )

    document.add_heading('6 测试集评估', level=1)
    metric_map = {
        'Precision': ('metrics/precision(B)', True),
        'Recall': ('metrics/recall(B)', True),
        'mAP50': ('metrics/mAP50(B)', True),
        'mAP50-95': ('metrics/mAP50-95(B)', True),
    }
    overall_rows = [[label, format_metric(overall.get(key), percent)] for label, (key, percent) in metric_map.items()]
    add_table(document, ['指标', '测试集结果'], overall_rows, '表 6-1 测试集总体指标')
    class_table_rows = []
    for item in per_class:
        class_table_rows.append(
            [
                item['class_name'],
                format_metric(item['precision'], True),
                format_metric(item['recall'], True),
                format_metric(item['mAP50'], True),
                format_metric(item['mAP50_95'], True),
            ]
        )
    add_table(document, ['类别', 'Precision', 'Recall', 'mAP50', 'mAP50-95'], class_table_rows, '表 6-2 测试集逐类别指标')
    add_image(document, EVAL_DIR / 'confusion_matrix_normalized.png', '图 6-1 测试集归一化混淆矩阵')

    document.add_heading('7 系统实现与演示', level=1)
    document.add_paragraph(
        '项目提供统一 CLI，包括 prepare、analyze、train、eval、predict、export 和 app 七个子命令。'
        '推理模块输出标注图片、JSON 汇总和 CSV 明细；Gradio 页面支持图片、摄像头与视频输入，并提供置信度、IoU 和输入尺寸调节。'
    )
    prediction = ASSETS / 'prediction_351.png'
    if prediction.exists():
        add_image(document, prediction, '图 7-1 测试样图检测结果', width_cm=12.5)
    web_demo = ASSETS / 'web_demo.png'
    if web_demo.exists():
        add_image(document, web_demo, '图 7-2 PCBA 缺陷检测网页界面', width_cm=9.5)
    document.add_paragraph('业务判定规则为：任一类别达到置信度阈值即判定为“不合格”，没有检测结果则判定为“合格”。')

    document.add_heading('8 局限与改进方向', level=1)
    document.add_paragraph(
        '当前实验在 CPU 上完成，输入尺寸和训练轮数受到时间限制。后续可以在 GPU 环境中提高分辨率、扩大 batch 和轮数，'
        '并比较多尺度训练、更多增广策略和模型融合。验证集与测试集来自不同编号区间，最佳验证 mAP50-95 为 36.98%，而测试集为 53.13%，表明当前固定编号划分可能包含批次或采集条件差异；后续应采用按批次分层或无泄漏的交叉验证划分。数据方面可增加真实产线的光照、遮挡和不同批次板卡样本，'
        '再结合误检样本回流形成持续迭代闭环。'
    )

    document.add_heading('附录 A 运行命令', level=1)
    commands = [
        r'pcba prepare --lighting auto --overwrite',
        r'pcba analyze',
        r'pcba train --config configs\train_cpu.yaml',
        r'pcba eval --split test --imgsz 960',
        r'pcba predict --source <图片或视频> --imgsz 960',
        r'pcba export --output models\pcba_yolov8n.onnx --imgsz 960',
        r'pcba app --host 127.0.0.1 --port 7860',
    ]
    for command in commands:
        paragraph = document.add_paragraph(style='No Spacing')
        run = paragraph.add_run(command)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
    return OUTPUT


if __name__ == '__main__':
    build_report()
